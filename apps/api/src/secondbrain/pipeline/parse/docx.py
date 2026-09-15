"""DOCX parser on python-docx.

Word documents carry explicit structure: paragraphs styled "Heading N" are
headings, "Title" is the document title, tables become one line per row.
Body and tables are walked in document order so offsets follow the page.
"""

import io
import re

import docx
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from secondbrain.pipeline.parse.base import ParsedDocument, ParseError, TextBuilder

_HEADING_STYLE = re.compile(r"^heading\s*(\d)", re.IGNORECASE)


def _iter_body(document: DocxDocument):
    """Paragraphs and tables in the order they appear in the body."""
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def parse_docx(data: bytes, *, fallback_title: str | None = None) -> ParsedDocument:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises PackageNotFoundError, KeyError, ValueError…
        raise ParseError(f"not a readable DOCX: {exc}") from exc

    builder = TextBuilder()
    title: str | None = None
    for item in _iter_body(document):
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue
            style = (item.style.name if item.style is not None else "") or ""
            if style.lower() == "title" and title is None:
                title = text
                builder.add_heading(1, text)
            elif m := _HEADING_STYLE.match(style):
                builder.add_heading(min(int(m.group(1)), 6), text)
            else:
                builder.add_paragraph(text)
        else:
            rows = []
            for row in item.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                builder.add_paragraph("\n".join(rows))

    text = builder.build()
    if not text.strip():
        raise ParseError("DOCX contains no text")

    core = document.core_properties
    title = title or (core.title or "").strip() or None
    if title is None and builder.headings:
        title = builder.headings[0].text
    meta = {
        "docx_properties": {
            k: v
            for k, v in {
                "author": core.author,
                "created": core.created.isoformat() if core.created else None,
                "modified": core.modified.isoformat() if core.modified else None,
                "subject": core.subject,
            }.items()
            if v
        }
    }
    return ParsedDocument(
        title=title or fallback_title, text=text, headings=builder.headings, meta=meta
    )
